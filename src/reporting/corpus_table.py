"""
Cumulative Corpus Table.

Joins together everything the pipeline knows about each paper into ONE table —
authors, venue, year, cluster, and every extracted field (problem/method/
results/inferences/novelty/limitations/datasets/metrics) — as one row per
paper. Outputs both a real .xlsx (for faculty who want to sort/filter in
Excel) and a self-contained .html report (open in any browser, no server
needed, click a row to expand full details) — this is the "read all papers
at once instead of opening one by one" view.

Deliberately joins against THREE existing files rather than requiring any
pipeline stage to be rerun:
  - collected_papers.json (Stage 1) — for authors, DOI, citation count. This
    is the fix for authors being silently dropped: Semantic Scholar returns
    them, but pdf_parser.py never copies that field forward into chunks, so
    it's lost by Stage 2 onward. Rather than plumbing "authors" through every
    intermediate file (chunks.json, paper_summaries.json, deep extractions),
    this script just re-joins against the one file that still has it.
  - paper_deep_extractions.json (Stage 5) if present, else paper_summaries.json
    (Stage 4) — whichever extraction is available, deep extraction preferred
    since it has the richer fields.
  - clusters.json — for the human-readable category label.

Run:
    python -m src.reporting.corpus_table --config configs/config.yaml
"""
import argparse
import html
import json
from pathlib import Path

import yaml


def load_config(config_path: str) -> dict:
    with open(config_path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def _authors_str(paper: dict) -> str:
    authors = paper.get("authors") or []
    names = [a.get("name", "") for a in authors if a.get("name")]
    return ", ".join(names) if names else ""


def build_corpus_table(config: dict) -> list[dict]:
    processed_dir = Path(config["paths"]["processed_dir"])
    raw_metadata_dir = Path(config["paths"]["raw_metadata_dir"])

    collected_path = raw_metadata_dir / "collected_papers.json"
    if not collected_path.exists():
        raise RuntimeError("collected_papers.json not found — run Stage 1 first.")
    collected = {p["paperId"]: p for p in json.loads(collected_path.read_text(encoding="utf-8"))}

    deep_path = processed_dir / "paper_deep_extractions.json"
    summaries_path = processed_dir / "paper_summaries.json"
    if deep_path.exists():
        extractions = {r["paper_id"]: r for r in json.loads(deep_path.read_text(encoding="utf-8"))}
        source_label = "deep extraction (Stage 5)"
    elif summaries_path.exists():
        extractions = {r["paper_id"]: r for r in json.loads(summaries_path.read_text(encoding="utf-8"))}
        source_label = "Stage 4 extraction (Stage 5 not run yet — fields will be shorter)"
    else:
        raise RuntimeError("No paper_summaries.json or paper_deep_extractions.json found — run Stage 4 first.")
    print(f"Using {source_label}")

    clusters = {}
    clusters_path = processed_dir / "clusters.json"
    if clusters_path.exists():
        raw_clusters = json.loads(clusters_path.read_text(encoding="utf-8"))
        clusters = {int(k): v for k, v in raw_clusters.items()}

    rows = []
    for paper_id, ex in extractions.items():
        collected_p = collected.get(paper_id, {})
        cluster_id = ex.get("cluster_id")
        cluster_label = clusters.get(cluster_id, {}).get("label", ex.get("category", ""))

        rows.append({
            "paper_id": paper_id,
            "title": ex.get("title") or collected_p.get("title", ""),
            "authors": _authors_str(collected_p),
            "year": ex.get("year") or collected_p.get("year", ""),
            "venue": ex.get("venue") or collected_p.get("venue", ""),
            "citation_count": collected_p.get("citationCount", ""),
            "doi": (collected_p.get("externalIds") or {}).get("DOI", ""),
            "category": cluster_label,
            "has_full_text": collected_p.get("has_full_text", False),
            "problem_addressed": ex.get("problem_addressed", ""),
            "method": ex.get("method", ""),
            "datasets": "; ".join(ex.get("datasets", []) or []),
            "metrics": "; ".join(ex.get("metrics", []) or []),
            "results": ex.get("results", ex.get("key_findings", "")),
            "inferences": ex.get("inferences", ""),
            "novelty_claim": ex.get("novelty_claim", ""),
            "limitations": ex.get("limitations", ""),
            "summary": ex.get("summary", ""),
            "extraction_flags": ", ".join(
                f for f in ["_extraction_failed", "_no_dataset_stated", "_dataset_fallback_used"]
                if ex.get(f)
            ),
        })

    rows.sort(key=lambda r: (r["category"], r["title"]))
    return rows


def write_xlsx(rows: list[dict], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Corpus"

    if not rows:
        ws.append(["No data"])
        wb.save(out_path)
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    header_fill = PatternFill(start_color="1C2430", end_color="1C2430", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    long_text_cols = {"problem_addressed", "method", "results", "inferences",
                       "novelty_claim", "limitations", "summary"}
    for row in rows:
        ws.append([row[h] for h in headers])

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 45 if header in long_text_cols else \
            (30 if header == "title" else 18)

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            header = headers[cell.column - 1]
            if header in long_text_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(out_path)


def write_html_report(rows: list[dict], out_path: Path, domain_query: str) -> None:
    def esc(v):
        return html.escape(str(v)) if v is not None else ""

    table_rows = []
    detail_panels = []
    for i, r in enumerate(rows):
        table_rows.append(f"""
        <tr onclick="toggle({i})">
          <td>{esc(r['title'])}</td>
          <td>{esc(r['authors'])}</td>
          <td>{esc(r['year'])}</td>
          <td>{esc(r['category'])}</td>
          <td>{esc(r['datasets'])}</td>
          <td>{'Yes' if r['has_full_text'] else 'No'}</td>
        </tr>""")

        flags_html = f"<p class='flags'>Flags: {esc(r['extraction_flags'])}</p>" if r['extraction_flags'] else ""
        detail_panels.append(f"""
        <tr id="detail-{i}" class="detail-row" style="display:none">
          <td colspan="6">
            <div class="detail">
              <h3>{esc(r['title'])}</h3>
              <p class="meta">{esc(r['authors'])} — {esc(r['venue'])}, {esc(r['year'])}
                 {f"· DOI: {esc(r['doi'])}" if r['doi'] else ""}
                 {f"· {esc(r['citation_count'])} citations" if r['citation_count'] != '' else ""}</p>
              {flags_html}
              <div class="field"><b>Problem addressed</b><p>{esc(r['problem_addressed']) or '—'}</p></div>
              <div class="field"><b>Method</b><p>{esc(r['method']) or '—'}</p></div>
              <div class="field"><b>Datasets</b><p>{esc(r['datasets']) or '—'}</p></div>
              <div class="field"><b>Metrics</b><p>{esc(r['metrics']) or '—'}</p></div>
              <div class="field"><b>Results</b><p>{esc(r['results']) or '—'}</p></div>
              <div class="field"><b>Inferences</b><p>{esc(r['inferences']) or '—'}</p></div>
              <div class="field"><b>Novelty claim</b><p>{esc(r['novelty_claim']) or '—'}</p></div>
              <div class="field"><b>Limitations</b><p>{esc(r['limitations']) or '—'}</p></div>
              <div class="field"><b>Summary</b><p>{esc(r['summary']) or '—'}</p></div>
            </div>
          </td>
        </tr>""")

    interleaved = "".join(t + d for t, d in zip(table_rows, detail_panels))

    html_doc = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Corpus Report</title>
<style>
  body {{ font-family: -apple-system, Inter, sans-serif; background:#e8ebe3; color:#1c2430; margin:0; padding:32px; }}
  h1 {{ font-size:22px; margin:0 0 4px; }}
  .subtitle {{ font-family: monospace; color:#57626e; margin:0 0 20px; }}
  table {{ width:100%; border-collapse:collapse; background:#f4f6f0; }}
  th {{ text-align:left; padding:10px 12px; background:#1c2430; color:white; font-size:12px;
        text-transform:uppercase; letter-spacing:.04em; position:sticky; top:0; }}
  td {{ padding:10px 12px; border-bottom:1px solid #d7ddd0; font-size:13px; }}
  tr:not(.detail-row):hover {{ background:#dfe4d8; cursor:pointer; }}
  .detail {{ padding:16px 8px; }}
  .detail h3 {{ margin:0 0 4px; font-size:17px; }}
  .meta {{ color:#57626e; font-size:12px; margin:0 0 12px; }}
  .flags {{ color:#b8842a; font-size:12px; }}
  .field {{ margin-bottom:12px; }}
  .field b {{ font-size:10.5px; text-transform:uppercase; letter-spacing:.06em; color:#b8842a; }}
  .field p {{ margin:4px 0 0; font-size:13.5px; line-height:1.5; }}
  input {{ padding:10px 14px; width:320px; margin-bottom:14px; border:1px solid #c1c9ba; border-radius:3px;
           font-family:monospace; }}
</style></head>
<body>
  <h1>Corpus Report</h1>
  <p class="subtitle">"{esc(domain_query)}" — {len(rows)} papers</p>
  <input id="search" placeholder="Filter by title, author, category..." onkeyup="filterRows()">
  <table id="corpus-table">
    <thead><tr><th>Title</th><th>Authors</th><th>Year</th><th>Category</th><th>Datasets</th><th>Full text</th></tr></thead>
    <tbody>{interleaved}</tbody>
  </table>
<script>
function toggle(i) {{
  const el = document.getElementById('detail-' + i);
  el.style.display = el.style.display === 'none' ? 'table-row' : 'none';
}}
function filterRows() {{
  const q = document.getElementById('search').value.toLowerCase();
  document.querySelectorAll('#corpus-table tbody tr:not(.detail-row)').forEach(tr => {{
    tr.style.display = tr.textContent.toLowerCase().includes(q) ? '' : 'none';
  }});
}}
</script>
</body></html>"""

    out_path.write_text(html_doc, encoding="utf-8")


def run_corpus_table(config: dict) -> None:
    processed_dir = Path(config["paths"]["processed_dir"])
    rows = build_corpus_table(config)
    print(f"Built table for {len(rows)} papers")

    xlsx_path = processed_dir / "corpus_table.xlsx"
    write_xlsx(rows, xlsx_path)
    print(f"Saved {xlsx_path}")

    html_path = processed_dir / "corpus_report.html"
    write_html_report(rows, html_path, config["collection"]["domain_query"])
    print(f"Saved {html_path} — open this directly in a browser, no server needed")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="configs/config.yaml")
    args = parser.parse_args()

    cfg = load_config(args.config)
    run_corpus_table(cfg)